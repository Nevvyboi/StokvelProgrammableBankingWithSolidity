#!/usr/bin/env python3
"""
Builds the PowerPoint decks for "The Treasurer Is a Smart Contract".

  python3 build.py            dark and light .pptx into ../
  python3 build.py --dark     one of them

Eighteen slides on the run sheet, two appendix slides for questions. Every slide is drawn from
code: real text runs, real shapes, the diagrams from assets/diagrams, speaker notes lifted from
talk/SPEECH.md. Fonts are referenced by name (Space Grotesk, Inter, JetBrains Mono); install
them from assets/fonts before opening, or use the PDF.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
ASSETS = ROOT / "assets"
sys.path.insert(0, str(ROOT / "cards"))
from generate import hash_seed, patch_cells, patch_colours  # noqa: E402

REPO_URL = "https://github.com/Nevvyboi/StokvelProgrammableBankingWithSolidity"
W, H = 13.333, 7.5
MX = 0.9

DISPLAY, BODY, MONO = "Space Grotesk", "Inter", "JetBrains Mono"

THEMES = {
    "dark": dict(bg="0B1426", surface="13203A", text="F4F1EA", muted="8A96AD", rule="2A3B5E", primary="2EC4B6", danger="FF5A5F", success="3DDC97", paper="F4F1EA", ink="0B1426", dim="4A5A7A"),
    "light": dict(bg="F7F4EC", surface="FFFFFF", text="0B1426", muted="5B6578", rule="C9CFDB", primary="0E9488", danger="D7333A", success="1B9E6A", paper="0B1426", ink="F4F1EA", dim="A9B0BE"),
}
MEMBERS = [("01", "Lerato", "F4A261"), ("02", "Thabo", "2A9D8F"), ("03", "Aisha", "E76F51"), ("04", "Johan", "8AB17D"), ("05", "Priya", "B388EB"), ("06", "Sipho", "E9C46A")]


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#").upper())


# ------------------------------------------------------------------ speaker notes from SPEECH.md
def load_notes() -> dict[int, str]:
    text = (ROOT / "talk" / "SPEECH.md").read_text(encoding="utf-8")
    parts = re.split(r"\n## (\d+) · ", text)
    notes: dict[int, str] = {}
    for i in range(1, len(parts), 2):
        n = int(parts[i])
        body = parts[i + 1].split("\n---")[0]
        body = "\n".join(body.split("\n")[1:]).strip()
        body = re.sub(r"\*([^*]+)\*", r"[\1]", body)  # stage directions in brackets
        body = re.sub(r"`([^`]+)`", r"\1", body)
        body = body.replace("★ ", "")
        notes[n] = body
    return notes


# ------------------------------------------------------------------ the deck
class Deck:
    def __init__(self, theme: str):
        self.t = THEMES[theme]
        self.theme = theme
        self.p = Presentation()
        self.p.slide_width = Inches(W)
        self.p.slide_height = Inches(H)
        self.notes = load_notes()
        self.n = 0

    # ---- primitives
    def slide(self, note_no: int | None = None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(self.t["bg"])
        self.n += 1
        if note_no and note_no in self.notes:
            s.notes_slide.notes_text_frame.text = self.notes[note_no]
        return s

    def text(self, s, x, y, w, h, runs, size=20, font=BODY, colour=None, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15, wrap=True):
        """runs: a string, or a list of paragraphs; each paragraph is a string or a list of (text, opts) runs."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        paragraphs = runs if isinstance(runs, list) else [runs]
        for i, para in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            pieces = para if isinstance(para, list) else [(para, {})]
            for piece, opts in pieces:
                r = p.add_run()
                r.text = piece
                f = r.font
                f.name = opts.get("font", font)
                f.size = Pt(opts.get("size", size))
                f.bold = opts.get("bold", bold)
                f.italic = opts.get("italic", False)
                f.color.rgb = rgb(opts.get("colour", colour or self.t["text"]))
        return tb

    def rule(self, s, x, y, w, colour=None, weight=1.0):
        ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = rgb(colour or self.t["rule"])
        ln.line.width = Pt(weight)
        return ln

    def rect(self, s, x, y, w, h, fill=None, line=None, weight=1.0, shape=MSO_SHAPE.RECTANGLE):
        r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            r.fill.solid()
            r.fill.fore_color.rgb = rgb(fill)
        else:
            r.fill.background()
        if line:
            r.line.color.rgb = rgb(line)
            r.line.width = Pt(weight)
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        return r

    def running_head(self, s, left, right="", y=0.62):
        self.rule(s, MX, y, W - 2 * MX, self.t["text"], 1.25)
        self.text(s, MX, y + 0.1, 8, 0.4, [[(left, {"font": DISPLAY, "bold": False, "size": 15})]], colour=self.t["text"])
        if right:
            self.text(s, W - MX - 7, y + 0.13, 7, 0.4, right, size=11, font=MONO, colour=self.t["muted"], align=PP_ALIGN.RIGHT)

    def footer(self, s, left, right="", size=11):
        self.rule(s, MX, H - 0.75, W - 2 * MX)
        self.text(s, MX, H - 0.65, (W - 2 * MX) if not right else 9, 0.3, left, size=size, font=MONO, colour=self.t["muted"])
        if right:
            self.text(s, W - MX - 4, H - 0.65, 4, 0.3, right, size=11, font=MONO, colour=self.t["muted"], align=PP_ALIGN.RIGHT)

    def stamp(self, s, x, y, label, colour, size=20, rot=-4.0):
        w = len(label) * size * 0.011 + 0.3
        h = size * 0.026
        outer = self.rect(s, x, y, w, h, line=colour, weight=2.25)
        inner = self.rect(s, x + 0.05, y + 0.05, w - 0.1, h - 0.1, line=colour, weight=0.75)
        tb = self.text(s, x, y, w, h, [[(label.upper(), {"font": DISPLAY, "bold": True, "size": size, "colour": colour})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for shp in (outer, inner, tb):
            shp.rotation = rot
        return tb

    def patch(self, s, x, y, size, colour, seed, base=None, n=4):
        """A quilt block, drawn with real shapes, from the same cells as the dashboard and the cards."""
        base = base or self.t["surface"]
        self.rect(s, x, y, size, size, fill=base)
        tones = patch_colours("#" + colour)
        c = size / n
        for (cx, cy, kind, rot, tone) in patch_cells(hash_seed(seed), n):
            if kind == "empty":
                continue
            px, py = x + cx * c, y + cy * c
            fill = tones[tone].lstrip("#")
            if kind == "square":
                self.rect(s, px, py, c, c, fill=fill)
            elif kind == "dot":
                self.rect(s, px + c / 6, py + c / 6, c * 2 / 3, c * 2 / 3, fill=fill, shape=MSO_SHAPE.OVAL)
            elif kind == "half":
                # rot 0: top half, 1: right, 2: bottom, 3: left
                if rot == 0:
                    self.rect(s, px, py, c, c / 2, fill=fill)
                elif rot == 1:
                    self.rect(s, px + c / 2, py, c / 2, c, fill=fill)
                elif rot == 2:
                    self.rect(s, px, py + c / 2, c, c / 2, fill=fill)
                else:
                    self.rect(s, px, py, c / 2, c, fill=fill)
            else:  # tri: right angle at one corner, rotating clockwise from top-left
                corners = [(px, py), (px + c, py), (px + c, py + c), (px, py + c)]
                a = corners[rot % 4]
                b = corners[(rot + 1) % 4]
                d = corners[(rot + 3) % 4]
                self.tri(s, a, b, d, fill)

    def tri(self, s, a, b, c_, fill):
        ff = s.shapes.build_freeform(Inches(a[0]), Inches(a[1]))
        ff.add_line_segments([(Inches(b[0]), Inches(b[1])), (Inches(c_[0]), Inches(c_[1]))], close=True)
        shp = ff.convert_to_shape()
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def picture(self, s, path, x, y, w=None, h=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)

    def diagram(self, s, name):
        p = ASSETS / "diagrams" / "png" / (f"{name}-light.png" if self.theme == "light" else f"{name}.png")
        return self.picture(s, p, 0, 0, w=W)

    def divider(self, s, n, title, promise, cut=False):
        self.running_head(s, f"Demo {n}", "live · sandbox and testnet only" + ("  ·  ✂ can be cut" if cut else ""))
        self.text(s, MX, 1.4, 6, 1.6, [[(f"DEMO {n}", {"font": DISPLAY, "size": 96, "bold": True, "colour": self.t["dim"]})]])
        self.text(s, MX, 3.15, 11.5, 1.2, [[(title, {"font": DISPLAY, "size": 54, "bold": False})]])
        self.text(s, MX, 4.45, 10.5, 1.3, promise, size=26, colour=self.t["muted"], spacing=1.25)
        # the quilt strip along the bottom edge
        for i, (num, _, col) in enumerate(MEMBERS):
            self.patch(s, MX + i * 0.62, H - 1.35, 0.55, col, f"member:{num}")

    # ---- code
    KEYWORDS = {"function", "if", "revert", "external", "view", "returns", "emit", "memory", "storage", "internal", "public", "pure", "mapping", "struct", "contract", "is", "override", "return", "uint256", "uint8", "uint128", "uint64", "bytes32", "bool", "address", "bytes", "calldata", "onlyRelayer", "immutable", "constant"}

    def code_runs(self, line: str, dim: bool):
        """Colour a Solidity line as runs. Dimmed lines are all one quiet colour."""
        if dim:
            return [(line, {"font": MONO, "size": 15, "colour": self.t["dim"]})]
        runs = []
        for tok in re.findall(r"//.*$|\"[^\"]*\"|[A-Za-z_][A-Za-z0-9_]*|\d+(?:_\d+)*|\s+|[^\sA-Za-z0-9_]+", line):
            if tok.startswith("//"):
                col = self.t["muted"]
            elif tok.startswith('"'):
                col = "E9C46A"
            elif tok in self.KEYWORDS:
                col = self.t["primary"]
            elif re.match(r"^\d", tok):
                col = "B388EB"
            elif tok[0].isupper() and len(tok) > 1 and not tok.isupper():
                col = "F4A261"  # errors and types
            else:
                col = self.t["text"]
            runs.append((tok, {"font": MONO, "size": 15, "colour": col}))
        return runs

    LINE = 0.0178  # inches per point of font size, one line at spacing 1.0

    def code(self, s, lines, highlight, x=MX, y=1.5, w=W - 2 * MX, size=15):
        paras = []
        for i, line in enumerate(lines):
            dim = i not in highlight
            runs = self.code_runs(line, dim)
            runs = [(t, {**o, "size": size}) for t, o in runs]
            paras.append(runs)
        lh = size * self.LINE
        self.rect(s, x - 0.25, y - 0.2, w + 0.5, len(lines) * lh + 0.4, fill=self.t["surface"], line=self.t["rule"])
        for i in highlight:
            self.rect(s, x - 0.25, y + i * lh, 0.06, lh, fill=self.t["primary"])
        self.text(s, x, y, w, len(lines) * lh + 0.2, paras, size=size, font=MONO, spacing=1.0, wrap=False)


# ------------------------------------------------------------------ slides
CLOSE_ROUND = """function closeRound() external {
    uint256 owed = pot + unsettled;
    if (reservesAt == 0 || reservesAt + RESERVES_MAX_AGE < block.timestamp) {
        revert ReservesStale(reservesAt, notBefore);
    }
    if (reservesCents < owed) revert ReservesShort(reservesCents, owed);
    if (block.timestamp < roundEndsAt) revert RoundStillOpen(roundEndsAt);

    uint8 recipient = uint8(round % members.length);
    uint256 cents = pot;
    payouts[round] = Payout(recipient, cents == 0, cents, members[recipient].beneficiaryHash);
    pot = 0;
    unsettled += cents;
    emit PayoutDue(round, recipient, cents, members[recipient].beneficiaryHash);

    round += 1;
    roundEndsAt = uint64(block.timestamp) + roundLength;
}""".split("\n")

RECORD_V2 = """function recordContribution(uint8 memberId, uint128 cents, bytes32 investecRef)
    external onlyRelayer
{
    if (seen[investecRef]) revert AlreadyRecorded(investecRef);
    if (memberId >= members.length) revert UnknownMember(memberId);
    if (cents > MAX_CONTRIBUTION_CENTS) revert AmountTooLarge(cents, MAX_CONTRIBUTION_CENTS);

    seen[investecRef] = true;
    members[memberId].paidCents += cents;
    pot += cents;
    totalIn += cents;
    emit ContributionRecorded(round, memberId, cents, investecRef);
}""".split("\n")


def build(theme: str, out: Path):
    d = Deck(theme)
    t = d.t

    # 1 title
    s = d.slide(1)
    d.text(s, MX, 0.9, 8, 0.4, "Investec developer community  ·  Q4 2026", size=13, font=MONO, colour=t["muted"])
    d.text(s, MX, 1.55, 9.4, 2.2, [[("The Treasurer", {"font": DISPLAY, "size": 58, "bold": True})], [("Is a Smart Contract", {"font": DISPLAY, "size": 58, "bold": True})]], spacing=1.0)
    d.text(s, MX, 3.85, 8.4, 0.9, "Rebuilding the stokvel with Solidity and Investec Programmable Banking", size=22, colour=t["text"])
    d.text(s, MX, 4.85, 8.4, 0.7, "Investec holds the rand. The chain holds the rulebook. The relayer carries messages and can't decide.", size=15, colour=t["muted"])
    d.text(s, MX, 6.45, 8, 0.4, [[("Nevin Tom", {"font": MONO, "size": 13, "bold": True}), ("   data engineer · MSc student · builds on Programmable Banking", {"font": MONO, "size": 13, "colour": t["muted"]})]])
    for i, (num, _, col) in enumerate(MEMBERS):
        d.patch(s, 10.6, 0.9 + i * 0.95, 0.9, col, f"member:{num}:hero")
        d.patch(s, 11.6, 0.9 + i * 0.95, 0.9, col, f"member:{num}")

    # 2 hook
    s = d.slide(2)
    d.running_head(s, "A group chat, December", "fictional · every name in this talk is made up")
    bubbles = [("Guys, where's December's money?", True), ("Guys?", True), ("…", False)]
    y = 1.6
    for i, (msg, face) in enumerate(bubbles):
        w = 7.0 if i == 0 else 2.2
        d.rect(s, MX, y, w, 0.9, fill=t["paper"], shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        d.text(s, MX + 0.3, y + 0.17, w - (1.4 if i == 0 else 0.6), 0.6, msg, size=24, colour=t["ink"], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        if i == 0:
            # a neutral face, drawn, so it renders everywhere
            fx = MX + w - 0.85
            d.rect(s, fx, y + 0.2, 0.5, 0.5, line=t["ink"], weight=2, shape=MSO_SHAPE.OVAL)
            d.rect(s, fx + 0.13, y + 0.36, 0.07, 0.07, fill=t["ink"], shape=MSO_SHAPE.OVAL)
            d.rect(s, fx + 0.30, y + 0.36, 0.07, 0.07, fill=t["ink"], shape=MSO_SHAPE.OVAL)
            d.rule(s, fx + 0.13, y + 0.55, 0.24, t["ink"], 2)
        d.text(s, MX + 0.1, y + 0.95, 3, 0.3, ["Priya · 09:41", "Priya · 09:58", "Thabo is typing"][i], size=11, font=MONO, colour=t["muted"])
        y += 1.45
    d.text(s, 8.6, 2.2, 4.0, 2.5, [[("Every stokvel has this thread.", {"font": DISPLAY, "size": 34})]], spacing=1.1)
    d.text(s, 8.6, 4.0, 4.0, 1.5, "One spreadsheet, one banking app, one person.", size=20, colour=t["muted"])

    # 3 scale
    s = d.slide(3)
    d.running_head(s, "Scale", "NASASA, August 2025")
    stats = [("11 million", "members"), ("810,000", "stokvels"), ("R50bn", "saved a year")]
    for i, (num, label) in enumerate(stats):
        x = MX + i * 3.95
        d.text(s, x, 2.0, 3.8, 1.4, [[(num, {"font": DISPLAY, "size": 66 if i else 60, "bold": True})]])
        d.text(s, x, 3.55, 3.8, 0.5, label, size=24, colour=t["muted"])
    d.text(s, MX, 5.1, 11.5, 0.8, "Fifty billion rand a year, moving on trust.", size=30, font=DISPLAY)
    d.footer(s, "Source: NASASA via EWN, 14 Aug 2025 · ewn.co.za/2025/08/14/more-than-11-million-south-africans-are-members-of-stokvels-nasasa", size=10)

    # 4 single point of trust
    s = d.slide(4)
    d.running_head(s, "Single point of trust", "how it works today")
    items = [("The ledger", "a notebook, a spreadsheet"), ("The money", "the treasurer's own account"), ("The rules", "the WhatsApp group")]
    for i, (a, b) in enumerate(items):
        y = 1.6 + i * 1.45
        d.rect(s, MX, y, 4.2, 1.1, fill=t["surface"], line=t["rule"])
        d.text(s, MX + 0.3, y + 0.18, 3.7, 0.5, [[(a, {"font": DISPLAY, "size": 24})]])
        d.text(s, MX + 0.3, y + 0.62, 3.7, 0.4, b, size=15, font=MONO, colour=t["muted"])
        ln = s.shapes.add_connector(1, Inches(MX + 4.2), Inches(y + 0.55), Inches(8.0), Inches(3.75))
        ln.line.color.rgb = rgb(t["muted"])
        ln.line.width = Pt(1.5)
    d.rect(s, 8.0, 2.75, 2.0, 2.0, line=t["text"], weight=2, shape=MSO_SHAPE.OVAL)
    d.text(s, 8.0, 2.75, 2.0, 2.0, [[("the", {"font": DISPLAY, "size": 16, "colour": t["muted"]})], [("treasurer", {"font": DISPLAY, "size": 24})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, 10.4, 2.9, 2.2, 1.8, [[("One person.", {"font": DISPLAY, "size": 22})], [("Three jobs.", {"font": DISPLAY, "size": 22})], [("No audit trail.", {"font": DISPLAY, "size": 22, "colour": t["danger"]})]], spacing=1.2)
    d.footer(s, "Investec doesn't sell a stokvel account. With Programmable Banking, it doesn't need to.")

    # 5 pattern
    s = d.slide(5)
    d.diagram(s, "pattern")

    # 6 hash chain
    s = d.slide(6)
    d.diagram(s, "hash-chain")

    # 7 demo 1
    s = d.slide(7)
    d.divider(s, 1, "The committee", "Six volunteers, six cards, one constitution. Nobody types a name in: the chain knows each member from a signature.")

    # 8 demo 2
    s = d.slide(8)
    d.divider(s, 2, "Swipe and round-up", "Aisha buys R87 of bread. The card rounds up to R90 and R3 lands on chain with her number on it.")

    # 9 code
    s = d.slide(9)
    d.running_head(s, "The whole payout rule", "contracts/src/Stokvel.sol")
    d.code(s, CLOSE_ROUND, highlight={2, 3, 4, 5, 6, 8}, y=1.35, size=14)
    d.text(s, MX, 6.35, 11.5, 0.6, [[("No function takes a recipient. ", {"font": DISPLAY, "size": 22}), ("round % 6", {"font": MONO, "size": 22, "colour": t["primary"]}), (", always.", {"font": DISPLAY, "size": 22})]])

    # 10 demo 3
    s = d.slide(10)
    d.divider(s, 3, "Month end", "The contract says who. Investec pays through the API. The chain records the bank's own payment reference.")

    # 11 demo 4
    s = d.slide(11)
    d.divider(s, 4, "The treasurer steals", "R1,000 leaves the account with the label “Admin fees”. Watch the payout freeze.")

    # 12 demo 5: the diff
    s = d.slide(12)
    d.running_head(s, "Demo 5 · Double webhook", "the one line between V1 and V2")
    d.code(s, RECORD_V2, highlight={3}, y=1.3, size=14)
    d.text(s, MX, 5.55, 11.5, 0.5, "V1 is this function without line 4. Same bank transaction, sent twice: V1 counts it twice, V2 says no.", size=20, colour=t["muted"])
    d.stamp(s, 10.2, 6.1, "refused", t["danger"], 24)

    # 13 demo 6 fuzzer (cut)
    s = d.slide(13)
    d.running_head(s, "Demo 6 · The fuzzer", "✂ can be cut for a 20 minute slot")
    d.text(s, MX, 1.3, 5.8, 0.6, [[("Four invariants, always", {"font": DISPLAY, "size": 26})]])
    inv = ["sum(paidCents) == totalIn", "totalIn - settledOut == pot + unsettled", "no Investec ref is ever used twice", "no close while reserves < owed"]
    d.text(s, MX, 2.05, 5.8, 2.6, [[(f"{i + 1}  ", {"font": MONO, "size": 15, "colour": t["primary"]}), (line, {"font": MONO, "size": 15})] for i, line in enumerate(inv)], spacing=1.6, wrap=False)
    d.text(s, MX, 4.6, 5.8, 1.2, "10,000 random months: contributions, duplicate webhooks, theft, refunds, time travel, closes, payouts.", size=17, colour=t["muted"])
    fail = ["$ ./scripts/fuzz.sh", "", "[FAIL: an Investec ref was recorded twice: 2 != 1]", "  contribute(uint8,uint128) args=[156, 2592000]", "  duplicateWebhook(uint256) args=[3]", "[FAIL: totalIn != what the bank received: 1184000 != 592000]", "", "Suite result: FAILED. 3 passed; 2 failed  (V1)", "Suite result: ok.     5 passed; 0 failed  (V2)"]
    d.rect(s, 7.0, 1.3, 5.45, 3.3, fill=t["surface"], line=t["rule"])
    paras = []
    for line in fail:
        col = t["danger"] if "FAIL" in line else t["success"] if "ok." in line else t["text"] if line.startswith("$") else t["muted"]
        paras.append([(line, {"font": MONO, "size": 12.5, "colour": col})])
    d.text(s, 7.25, 1.5, 5.1, 3.0, paras, spacing=1.3, wrap=False)
    d.text(s, MX, 6.35, 11.5, 0.5, [[("Write down what must always be true. Let a machine try to make it false.", {"font": DISPLAY, "size": 22})]])

    # 14 demo 7
    s = d.slide(14)
    d.divider(s, 7, "Vote", "Four cards sign one English sentence: “raise the contribution from R5,000 to R6,000”. Then a timelock. Then it's the rule.")

    # 15 demo 8
    s = d.slide(15)
    d.divider(s, 8, "Stolen card", "The thief holds Sipho's key. The committee moves Sipho to a spare card, and the stolen one stops being anyone.")

    # 16 kill the server (cut)
    s = d.slide(16)
    d.divider(s, 9, "Kill the server", "The relayer dies on stage. The book is still on the shelf: every screen rebuilds from chain events, and nothing is sent twice.", cut=True)

    # 17 honest limits
    s = d.slide(17)
    d.running_head(s, "Honest limits, and what's next", "said out loud")
    cols = [("What this doesn't do", [("Trusted relayer", "it reports what the bank did"), ("Group account KYC", "six people on one key is a demo"), ("No tokenised rand", "yet, on purpose"), ("QR cards", "should be passkeys")], t["danger"]),
            ("Next", [("Passkeys", "a phone that signs"), ("A real club account", "with the bank's own controls"), ("Rage-quit and late fees", "the rules a real stokvel has"), ("Tokenised rand", "when a bank stands behind one")], t["primary"])]
    for ci, (head, items, col) in enumerate(cols):
        x = MX + ci * 6.0
        d.text(s, x, 1.35, 5.5, 0.5, [[(head, {"font": DISPLAY, "size": 26, "colour": col})]])
        for i, (a, b) in enumerate(items):
            y = 2.05 + i * 0.95
            d.rule(s, x, y, 5.4)
            d.text(s, x, y + 0.12, 5.4, 0.4, [[(a, {"font": DISPLAY, "size": 22})]])
            d.text(s, x, y + 0.5, 5.4, 0.35, b, size=15, colour=t["muted"])
    d.footer(s, "With thanks to Peter Smythe · github.com/petersmythe/invapi-dual-auth · github.com/petersmythe/investec-swipe-n-save", size=10)

    # 18 close
    s = d.slide(18)
    d.text(s, MX, 1.3, 8.4, 2.9, [[("We didn't replace the trust.", {"font": DISPLAY, "size": 44})], [("We wrote it down where nobody can erase it.", {"font": DISPLAY, "size": 44, "colour": t["primary"]})]], spacing=1.15)
    d.text(s, MX, 4.6, 8.4, 0.5, "Thank you, committee.", size=20, colour=t["muted"])
    for i, (num, name, col) in enumerate(MEMBERS):
        d.patch(s, MX + i * 1.35, 5.15, 0.55, col, f"member:{num}")
        d.text(s, MX + i * 1.35, 5.78, 1.3, 0.3, name, size=13, font=MONO, colour=t["muted"])
    qr_path = HERE / f"qr-{theme}.png"
    img = qrcode.make(REPO_URL, box_size=10, border=1)
    img.save(qr_path)
    d.rect(s, 9.95, 1.4, 2.5, 2.5, fill="F4F1EA")
    d.picture(s, qr_path, 10.05, 1.5, w=2.3)
    d.text(s, 9.2, 4.05, 4.0, 0.9, [[("github.com/Nevvyboi/", {"font": MONO, "size": 12, "colour": t["muted"]})], [("StokvelProgrammableBankingWithSolidity", {"font": MONO, "size": 12, "colour": t["muted"]})]], align=PP_ALIGN.CENTER)
    d.footer(s, "one command runs the whole demo: ./scripts/demo.sh --mock", "MIT")

    # 19 appendix: threat model
    s = d.slide()
    d.running_head(s, "Appendix · Threat model", "for questions")
    rows = [("Dodgy treasurer", "\u201cpay me this month\u201d", "no recipient setter; vote plus timelock", "no setter", t["success"]),
            ("Dodgy treasurer", "takes rand out of the account", "ReservesShort freezes payouts", "frozen", t["danger"]),
            ("Compromised relayer", "records a payment twice", "seen[ref]: AlreadyRecorded", "refused", t["danger"]),
            ("Compromised relayer", "invents a huge contribution", "AmountTooLarge cap; short at close", "capped", t["success"]),
            ("Compromised relayer", "pays the wrong person", "beneficiaryHash check; vote it out", "checked", t["success"]),
            ("Vote thief", "replays an old signed vote", "EIP-712 domain, per-signer nonce", "nonce", t["success"]),
            ("Stolen member card", "votes as the victim", "RotateKey to a spare card", "rotated", t["success"])]
    for i, (who, attack, defence, st, col) in enumerate(rows):
        y = 1.3 + i * 0.72
        d.rule(s, MX, y + 0.62, W - 2 * MX)
        d.text(s, MX, y + 0.1, 2.5, 0.45, [[(who, {"font": DISPLAY, "size": 16})]])
        d.text(s, 3.5, y + 0.1, 3.1, 0.45, attack, size=15)
        d.text(s, 6.7, y + 0.1, 4.3, 0.45, defence, size=13.5, colour=t["muted"])
        d.stamp(s, 11.0, y + 0.08, st, col, 12, -3)
    d.footer(s, "docs/SECURITY.md has the full model, the invariants and the Slither results")

    # 20 appendix: architecture
    s = d.slide()
    d.diagram(s, "architecture")

    d.p.save(out)
    qr_path.unlink(missing_ok=True)
    print(f"wrote {out} ({d.n} slides)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dark", action="store_true")
    ap.add_argument("--light", action="store_true")
    a = ap.parse_args()
    both = not (a.dark or a.light)
    if a.dark or both:
        build("dark", HERE.parent / "The-Treasurer-Is-a-Smart-Contract.pptx")
    if a.light or both:
        build("light", HERE.parent / "The-Treasurer-Is-a-Smart-Contract-Light.pptx")
